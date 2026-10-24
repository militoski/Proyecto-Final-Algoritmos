import os
import smtplib
import ssl
from email.message import EmailMessage
from dotenv import load_dotenv

load_dotenv()

def enviar_correo(destinatario, asunto, mensaje, archivo_adjunto=None):
    correo = os.getenv("GMAIL_CORREO")
    clave = os.getenv("GMAIL_CLAVE")

    if not correo or not clave:
        print("Error: revisa los datos del archivo .env")
        return False

    email = EmailMessage()
    email["From"] = correo
    email["To"] = destinatario
    email["Subject"] = asunto
    email.set_content(mensaje)

    if archivo_adjunto:
        with open(archivo_adjunto, "rb") as archivo:
            email.add_attachment(
                archivo.read(),
                maintype="application",
                subtype="vnd.openxmlformats-officedocument.wordprocessingml.document",
                filename="cotizacion.docx")
    try:
        contexto = ssl.create_default_context()

        with smtplib.SMTP_SSL(
            "smtp.gmail.com", 465, context=contexto) as servidor:
            servidor.login(correo, clave)
            servidor.send_message(email)

        print("Correo enviado correctamente.")
        return True

    except Exception as error:
        print("No se pudo enviar el correo:", error)
        return False

def enviar_copia_seguridad(destinatario):
    correo = os.getenv("GMAIL_CORREO")
    clave = os.getenv("GMAIL_CLAVE")

    email = EmailMessage()
    email["From"] = correo
    email["To"] = destinatario
    email["Subject"] = "Copia de seguridad del proyecto"
    email.set_content(
        "Adjunto las copias de seguridad del inventario, "
        "clientes y pedidos del sistema."
    )

    try:
        archivos = [
            "copia_inventario.txt",
            "copia_clientes.txt",
            "copia_pedidos.txt"
        ]

        for nombre_archivo in archivos:
            with open(nombre_archivo, "rb") as archivo:
                email.add_attachment(
                    archivo.read(),
                    maintype="text",
                    subtype="plain",
                    filename=nombre_archivo
                )

        contexto = ssl.create_default_context()

        with smtplib.SMTP_SSL(
            "smtp.gmail.com", 465, context=contexto
        ) as servidor:
            servidor.login(correo, clave)
            servidor.send_message(email)

        print("Copia de seguridad enviada correctamente.")

    except Exception as error:
        print("Error al enviar la copia:", error)

opcion = ""

productos = []
clientes = []
pedidos = []

# Cargar productos guardadso
try:
    archivo = open("inventario.txt", "r")

    for linea in archivo:
        datos = linea.strip().split(",")

        if len(datos) == 3:
            producto = {"nombre": datos[0],
                "precio": float(datos[1]),
                "existencia": int(datos[2])}

            productos.append(producto)

    archivo.close()

except FileNotFoundError:
    archivo = open("inventario.txt", "w")
    archivo.close()


# Cargar clientes guardados
try:
    archivo = open("clientes.txt", "r")

    for linea in archivo:
        datos = linea.strip().split(",")

        if len(datos) == 3:
            cliente = {
                "nombre": datos[0],
                "nit": datos[1],
                "direccion": datos[2]}

            clientes.append(cliente)

    archivo.close()

except FileNotFoundError:
    archivo = open("clientes.txt", "w")
    archivo.close()


# Cargar pedidos guardados
try:
    archivo = open("pedidos.txt", "r")

    for linea in archivo:
        datos = linea.strip().split(",")

        if len(datos) == 4:
            pedido = {"cliente": datos[0],
                "producto": datos[1],
                "cantidad": int(datos[2]),
                "valor": float(datos[3])}

            pedidos.append(pedido)

    archivo.close()

except FileNotFoundError:
    archivo = open("pedidos.txt", "w")
    archivo.close()

while opcion != "6":
    print("\n==============================")
    print("    SISTEMA DE INVENTARIO")
    print("==============================")
    print("1. Productos")
    print("2. Clientes")
    print("3. Pedidos")
    print("4. Informes")
    print("5. Varios")
    print("6. Salir")

    opcion = input("Seleccione una opcion: ")

    if opcion == "1":
        print("\n=== PRODUCTOS ===")
        print("1. Agregar producto")
        print("2. Listar productos")
        print("3. Editar producto")
        print("4. Eliminar producto")
        print("5. Crear cotizacion")
        opcion_producto = input("Seleccione una opcion: ")

        if opcion_producto == "1":
            nombre = input("Ingrese el nombre del producto: ")
            precio = float(input("Ingrese el precio: "))
            existencia = int(input("Ingrese la existencia: "))

            producto = {"nombre": nombre,
                "precio": precio,
                "existencia": existencia}

            productos.append(producto)

            archivo = open("inventario.txt", "a")
            archivo.write(
                nombre + "," + str(precio) + ","
                + str(existencia) + "\n")
            archivo.close()

            print("Producto agregado correctamente")

        elif opcion_producto == "2":
            print("\nLISTA DE PRODUCTOS")

            if len(productos) == 0:
                print("No hay productos registrados")
            else:
                for producto in productos:
                    print("Nombre:", producto["nombre"])
                    print("Precio:", producto["precio"])
                    print("Existencia:", producto["existencia"])
                    print("----------------------")

        elif opcion_producto == "3":
            nombre_buscar = input(
                "Ingrese el nombre del producto a editar: ")
            encontrado = False

            for producto in productos:
                if producto["nombre"].lower() == nombre_buscar.lower():
                    print("Producto encontrado")

                    nuevo_nombre = input("Ingrese el nuevo nombre: ")
                    nuevo_precio = float(input("Ingrese el nuevo precio: "))
                    nueva_existencia = int(
                        input("Ingrese la nueva existencia: "))

                    producto["nombre"] = nuevo_nombre
                    producto["precio"] = nuevo_precio
                    producto["existencia"] = nueva_existencia

                    encontrado = True
                    break

            if encontrado:
                archivo = open("inventario.txt", "w")

                for producto in productos:
                    archivo.write(
                        producto["nombre"] + ","
                        + str(producto["precio"]) + ","
                        + str(producto["existencia"]) + "\n")

                archivo.close()
                print("Producto editado correctamente")
            else:
                print("Producto no encontrado")

        elif opcion_producto == "4":
            nombre_buscar = input(
                "Ingrese el nombre del producto a eliminar: ")
            encontrado = False

            for producto in productos:
                if producto["nombre"].lower() == nombre_buscar.lower():
                    productos.remove(producto)
                    encontrado = True
                    break

            if encontrado:
                archivo = open("inventario.txt", "w")

                for producto in productos:
                    archivo.write(
                        producto["nombre"] + ","
                        + str(producto["precio"]) + ","
                        + str(producto["existencia"]) + "\n")

                archivo.close()
                print("Producto eliminado correctamente")
            else:
                print("Producto no encontrado")

        elif opcion_producto == "5":
            from docx import Document

            nombre_cliente = input("Nombre del cliente: ")
            correo_cliente = input("Correo del cliente: ")
            nombre_producto = input("Producto que desea cotizar: ")

            encontrado = False

            for producto in productos:
                if producto["nombre"].lower() == nombre_producto.lower():
                    documento = Document()

                    documento.add_heading("COTIZACION", level=1)
                    documento.add_paragraph(
                        "Estimado " + nombre_cliente
                    )
                    documento.add_paragraph(
                        "Correo: " + correo_cliente
                    )
                    documento.add_paragraph(
                        "El precio de nuestro producto "
                        + producto["nombre"]
                    )
                    documento.add_paragraph(
                        "es Q" + str(producto["precio"])
                    )

                    documento.save("cotizacion.docx")

                    print("Cotizacion creada correctamente")
                    print("Archivo: cotizacion.docx")
                    enviar_correo(correo_cliente,"Cotizacion de nuestro producto","Estimado "
             + nombre_cliente + ", adjuntamos la cotizacion solicitada.","cotizacion.docx")
                    encontrado = True
                    break

            if encontrado == False:
                print("Producto no encontrado")

        else:
            print("Opcion incorrecta")

    
    elif opcion == "2":
        print("\n=== CLIENTES ===")
        print("1. Agregar cliente")
        print("2. Listar clientes")
        print("3. Editar cliente")
        print("4. Eliminar cliente")

        opcion_cliente = input("Seleccione una opcion: ")

        if opcion_cliente == "1":
            nombre = input("Ingrese el nombre del cliente: ")
            nit = input("Ingrese el NIT del cliente: ")
            direccion = input("Ingrese la direccion: ")

            cliente = {"nombre": nombre,
                "nit": nit,
                "direccion": direccion}

            clientes.append(cliente)

            archivo = open("clientes.txt", "a")
            archivo.write(nombre + "," + nit + "," + direccion + "\n")
            archivo.close()

            print("Cliente agregado correctamente")

        elif opcion_cliente == "2":
            print("\nLISTA DE CLIENTES")

            if len(clientes) == 0:
                print("No hay clientes registrados")
            else:
                for cliente in clientes:
                    print("Nombre:", cliente["nombre"])
                    print("NIT:", cliente["nit"])
                    print("Direccion:", cliente["direccion"])
                    print("----------------------")

        elif opcion_cliente == "3":
            nombre_buscar = input("Ingrese el nombre del cliente a editar: ")
            encontrado = False

            for cliente in clientes:
                if cliente["nombre"].lower() == nombre_buscar.lower():
                    print("Cliente encontrado")

                    nuevo_nombre = input("Ingrese el nuevo nombre: ")
                    nuevo_nit = input("Ingrese el nuevo NIT: ")
                    nueva_direccion = input("Ingrese la nueva direccion: ")

                    cliente["nombre"] = nuevo_nombre
                    cliente["nit"] = nuevo_nit
                    cliente["direccion"] = nueva_direccion

                    encontrado = True
                    break

            if encontrado:
                archivo = open("clientes.txt", "w")

                for cliente in clientes:
                    archivo.write(
                        cliente["nombre"] + ","
                        + cliente["nit"] + ","
                        + cliente["direccion"] + "\n")

                archivo.close()
                print("Cliente editado correctamente")
            else:
                print("Cliente no encontrado")

        elif opcion_cliente == "4":
            nombre_buscar = input("Ingrese el nombre del cliente a eliminar: ")
            encontrado = False

            for cliente in clientes:
                if cliente["nombre"].lower() == nombre_buscar.lower():
                    clientes.remove(cliente)
                    encontrado = True
                    break

            if encontrado:
                archivo = open("clientes.txt", "w")

                for cliente in clientes:
                    archivo.write(
                        cliente["nombre"] + ","
                        + cliente["nit"] + ","
                        + cliente["direccion"] + "\n")

                archivo.close()
                print("Cliente eliminado correctamente")
            else:
                print("Cliente no encontrado")

        else:
            print("Opcion incorrecta")


    
    elif opcion == "3":
        print("\n=== PEDIDOS ===")
        print("1. Agregar pedido")
        print("2. Listar pedidos")
        print("3. Eliminar pedido")

        opcion_pedido = input("Seleccione una opcion: ")

        if opcion_pedido == "1":
            nombre_cliente = input("Ingrese el nombre del cliente: ")
            nombre_producto = input("Ingrese el nombre del producto: ")

            cantidad = int(input("Ingrese la cantidad: "))

            producto_encontrado = False
            valor = 0

            for producto in productos:
                if producto["nombre"].lower() == nombre_producto.lower():
                    producto_encontrado = True
                    valor = producto["precio"] * cantidad
                    break

            if producto_encontrado and cantidad > 0:
                pedido = {
                    "cliente": nombre_cliente,
                    "producto": nombre_producto,
                    "cantidad": cantidad,
                    "valor": valor
                }

                pedidos.append(pedido)

                archivo = open("pedidos.txt", "a")
                archivo.write(
                    nombre_cliente + ","
                    + nombre_producto + ","
                    + str(cantidad) + ","
                    + str(valor) + "\n")
                archivo.close()

                print("Pedido agregado correctamente")
                print("Valor del pedido:", valor)

            elif not producto_encontrado:
                print("El producto no existe")
            else:
                print("La cantidad debe ser mayor que cero")

        elif opcion_pedido == "2":
            print("\nLISTA DE PEDIDOS")

            if len(pedidos) == 0:
                print("No hay pedidos registrados")
            else:
                numero = 1

                for pedido in pedidos:
                    print("Pedido:", numero)
                    print("Cliente:", pedido["cliente"])
                    print("Producto:", pedido["producto"])
                    print("Cantidad:", pedido["cantidad"])
                    print("Valor:", pedido["valor"])
                    print("----------------------")

                    numero = numero + 1

        elif opcion_pedido == "3":
            if len(pedidos) == 0:
                print("No hay pedidos para eliminar")
            else:
                print("\nLISTA DE PEDIDOS")

                numero = 1

                for pedido in pedidos:
                    print(
                        numero, "-", pedido["cliente"],
                        "-", pedido["producto"],
                        "-", pedido["valor"])
                    numero = numero + 1

                numero_eliminar = int(
                    input("Ingrese el numero del pedido a eliminar: "))

                if numero_eliminar >= 1 and numero_eliminar <= len(pedidos):
                    pedidos.pop(numero_eliminar - 1)

                    archivo = open("pedidos.txt", "w")

                    for pedido in pedidos:
                        archivo.write(
                            pedido["cliente"] + ","
                            + pedido["producto"] + ","
                            + str(pedido["cantidad"]) + ","
                            + str(pedido["valor"]) + "\n")

                    archivo.close()

                    print("Pedido eliminado correctamente")
                else:
                    print("Numero de pedido incorrecto")

        else:
            print("Opcion incorrecta")


   
    elif opcion == "4":
        print("\n=== INFORMES ===")
        print("1. Total de venta por cliente")
        print("2. Total de ventas por producto")

        opcion_informe = input("Seleccione una opcion: ")

        if opcion_informe == "1":
            nombre_cliente = input("Ingrese el nombre del cliente: ")
            total = 0
            encontrado = False

            for pedido in pedidos:
                if pedido["cliente"].lower() == nombre_cliente.lower():
                    total = total + pedido["valor"]
                    encontrado = True

            if encontrado:
                print("Cliente:", nombre_cliente)
                print("Total de ventas:", round(total, 2))
            else:
                print("No se encontraron ventas para ese cliente")

        elif opcion_informe == "2":
            nombre_producto = input("Ingrese el nombre del producto: ")
            total = 0
            encontrado = False

            for pedido in pedidos:
                if pedido["producto"].lower() == nombre_producto.lower():
                    total = total + pedido["valor"]
                    encontrado = True

            if encontrado:
                print("Producto:", nombre_producto)
                print("Total de ventas:", round(total, 2))
            else:
                print("No se encontraron ventas para ese producto")

        else:
            print("Opcion incorrecta")


    
    elif opcion == "5":
        print("\n=== VARIOS ===")
        print("1. Crear copia de seguridad")

        opcion_varios = input("Seleccione una opcion: ")

        if opcion_varios == "1":
            archivo = open("inventario.txt", "r")
            copia = open("copia_inventario.txt", "w")
            copia.write(archivo.read())
            archivo.close()
            copia.close()

            archivo = open("clientes.txt", "r")
            copia = open("copia_clientes.txt", "w")
            copia.write(archivo.read())
            archivo.close()
            copia.close()

            archivo = open("pedidos.txt", "r")
            copia = open("copia_pedidos.txt", "w")
            copia.write(archivo.read())
            archivo.close()
            copia.close()

            print("Copia de seguridad creada correctamente")

            correo_coordinador = input(
                "Ingrese el correo del coordinador: "
            )

            enviar_copia_seguridad(correo_coordinador)
